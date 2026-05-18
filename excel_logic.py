import io
import re
from copy import copy

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import MultiCellRange

CLASS_NAME_PATTERN = re.compile(r"^\d+[A-Z]\d+$")


def extract_class_from_filename(filename):
    """Trich xuat ten lop tu ten file. VD: 4A1_HUYNH_MINH_DANG.xlsx -> 4A1"""
    stem = filename.rsplit(".", 1)[0]
    first_part = stem.split("_")[0].strip().upper()
    if CLASS_NAME_PATTERN.fullmatch(first_part):
        return first_part, None
    return None, "Không xác định được tên lớp (phần trước '_' phải dạng xAy, VD: 4A1)"


def copy_cell_style(style_data, target_cell):
    if style_data.get("font") is not None:
        target_cell.font = copy(style_data["font"])
    if style_data.get("border") is not None:
        target_cell.border = copy(style_data["border"])
    if style_data.get("fill") is not None:
        target_cell.fill = copy(style_data["fill"])
    if style_data.get("alignment") is not None:
        target_cell.alignment = copy(style_data["alignment"])
    if style_data.get("number_format") is not None:
        target_cell.number_format = style_data["number_format"]


def extract_cell_payload(cell):
    return {
        "value": cell.value,
        "font": copy(cell.font) if cell.font is not None else None,
        "border": copy(cell.border) if cell.border is not None else None,
        "fill": copy(cell.fill) if cell.fill is not None else None,
        "number_format": cell.number_format,
        "alignment": copy(cell.alignment) if cell.alignment is not None else None,
    }


def to_range_string(min_col, min_row, max_col, max_row):
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def get_merged_boundaries(ws):
    return [range_boundaries(str(rng)) for rng in ws.merged_cells.ranges]


def remap_merged_ranges_for_insert(merged_ranges, start_row, amount):
    remapped = []
    for min_col, min_row, max_col, max_row in merged_ranges:
        if max_row < start_row:
            remapped.append((min_col, min_row, max_col, max_row))
        elif min_row >= start_row:
            remapped.append((min_col, min_row + amount, max_col, max_row + amount))
        else:
            remapped.append((min_col, min_row, max_col, max_row + amount))
    return remapped


def remap_merged_ranges_for_delete(merged_ranges, start_row, amount):
    end_row = start_row + amount - 1
    remapped = []

    for min_col, min_row, max_col, max_row in merged_ranges:
        if max_row < start_row:
            remapped.append((min_col, min_row, max_col, max_row))
            continue

        if min_row > end_row:
            remapped.append((min_col, min_row - amount, max_col, max_row - amount))
            continue

        if min_row >= start_row and max_row <= end_row:
            continue

        if min_row < start_row and max_row > end_row:
            remapped.append((min_col, min_row, max_col, max_row - amount))
            continue

        if min_row < start_row <= max_row <= end_row:
            new_max_row = start_row - 1
            if new_max_row >= min_row:
                remapped.append((min_col, min_row, max_col, new_max_row))
            continue

        if start_row <= min_row <= end_row < max_row:
            new_min_row = start_row
            new_max_row = max_row - amount
            if new_max_row >= new_min_row:
                remapped.append((min_col, new_min_row, max_col, new_max_row))

    return remapped


def apply_merged_ranges(ws, merged_ranges):
    ws._cells = {
        coord: cell
        for coord, cell in ws._cells.items()
        if not isinstance(cell, MergedCell)
    }
    ws.merged_cells = MultiCellRange()

    seen = set()
    for min_col, min_row, max_col, max_row in merged_ranges:
        if min_row > max_row or min_col > max_col:
            continue
        range_str = to_range_string(min_col, min_row, max_col, max_row)
        if range_str in seen:
            continue
        ws.merge_cells(range_str)
        seen.add(range_str)


def snapshot_row_heights(ws, start_row):
    return {
        row_idx: dim.height
        for row_idx, dim in ws.row_dimensions.items()
        if row_idx >= start_row and dim.height is not None
    }


def shift_row_heights_after_insert(ws, row_heights, start_row, amount):
    for row_idx in sorted(row_heights.keys(), reverse=True):
        ws.row_dimensions[row_idx + amount].height = row_heights[row_idx]

    for row_idx in range(start_row, start_row + amount):
        ws.row_dimensions[row_idx].height = None


def _to_bytes_buffer(uploaded_file):
    if hasattr(uploaded_file, "getvalue"):
        return io.BytesIO(uploaded_file.getvalue())

    if hasattr(uploaded_file, "read"):
        data = uploaded_file.read()
        if hasattr(uploaded_file, "seek"):
            uploaded_file.seek(0)
        return io.BytesIO(data)

    if isinstance(uploaded_file, (bytes, bytearray)):
        return io.BytesIO(uploaded_file)

    raise TypeError("Unsupported uploaded_file type")


def process_excel(uploaded_file, class_name, filename):
    wb = None
    try:
        source_start_row = 20
        cut_amount = 8
        source_end_row = source_start_row + cut_amount - 1
        target_insert_row = 2

        file_buffer = _to_bytes_buffer(uploaded_file)
        wb = openpyxl.load_workbook(file_buffer)

        src_sheet_name = f"MH {class_name} (2025-2026)"
        tgt_sheet_name = f"NL, PC {class_name} (2025-2026)"

        if src_sheet_name not in wb.sheetnames or tgt_sheet_name not in wb.sheetnames:
            missing_sheets = [
                sheet_name
                for sheet_name in (src_sheet_name, tgt_sheet_name)
                if sheet_name not in wb.sheetnames
            ]
            return None, f"❌ File '{filename}' thiếu sheet: {', '.join(missing_sheets)}."

        ws_src = wb[src_sheet_name]
        ws_tgt = wb[tgt_sheet_name]

        src_merged_ranges = get_merged_boundaries(ws_src)
        cut_block_merged_ranges = []
        for min_col, min_row, max_col, max_row in src_merged_ranges:
            overlaps = not (max_row < source_start_row or min_row > source_end_row)
            fully_inside = min_row >= source_start_row and max_row <= source_end_row

            if overlaps and not fully_inside:
                range_str = to_range_string(min_col, min_row, max_col, max_row)
                return (
                    None,
                    f"❌ File '{filename}' có merge cell cắt ngang vùng dòng 20-27 ({range_str}), không thể cắt an toàn.",
                )

            if fully_inside:
                new_min_row = target_insert_row + (min_row - source_start_row)
                new_max_row = target_insert_row + (max_row - source_start_row)
                cut_block_merged_ranges.append((min_col, new_min_row, max_col, new_max_row))

        rows_data = []
        source_row_heights = []
        for row_idx, row in enumerate(
            ws_src.iter_rows(
                min_row=source_start_row,
                max_row=source_end_row,
                max_col=ws_src.max_column,
            ),
            start=source_start_row,
        ):
            row_data = [extract_cell_payload(cell) for cell in row]
            rows_data.append(row_data)
            source_row_heights.append(ws_src.row_dimensions[row_idx].height)

        tgt_merged_ranges_before_insert = get_merged_boundaries(ws_tgt)
        tgt_row_heights_before_insert = snapshot_row_heights(ws_tgt, start_row=target_insert_row)

        ws_tgt.insert_rows(target_insert_row, amount=cut_amount)

        remapped_tgt_merged_ranges = remap_merged_ranges_for_insert(
            tgt_merged_ranges_before_insert,
            start_row=target_insert_row,
            amount=cut_amount,
        )
        apply_merged_ranges(ws_tgt, remapped_tgt_merged_ranges)
        shift_row_heights_after_insert(
            ws_tgt,
            tgt_row_heights_before_insert,
            start_row=target_insert_row,
            amount=cut_amount,
        )

        for r_offset, row_data in enumerate(rows_data):
            target_row_idx = target_insert_row + r_offset
            row_height = source_row_heights[r_offset]
            if row_height is not None:
                ws_tgt.row_dimensions[target_row_idx].height = row_height

            for c_idx, cell_data in enumerate(row_data, start=1):
                tgt_cell = ws_tgt.cell(row=target_row_idx, column=c_idx)
                tgt_cell.value = cell_data["value"]
                copy_cell_style(cell_data, tgt_cell)

        for min_col, min_row, max_col, max_row in cut_block_merged_ranges:
            ws_tgt.merge_cells(to_range_string(min_col, min_row, max_col, max_row))

        ws_src.delete_rows(source_start_row, amount=cut_amount)
        remapped_src_merged_ranges = remap_merged_ranges_for_delete(
            src_merged_ranges,
            start_row=source_start_row,
            amount=cut_amount,
        )
        apply_merged_ranges(ws_src, remapped_src_merged_ranges)

        out_stream = io.BytesIO()
        wb.save(out_stream)
        out_stream.seek(0)

        return out_stream, None
    except Exception as e:
        return None, f"❌ Lỗi khi xử lý file '{filename}': {str(e)}"
    finally:
        if wb is not None:
            wb.close()
