import streamlit as st
import openpyxl
from copy import copy
import io
import zipfile
import re
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.cell_range import MultiCellRange

# Cấu hình trang UI
st.set_page_config(page_title="Công cụ Xử lý Excel Dữ liệu Lớp", layout="centered")

CLASS_NAME_PATTERN = re.compile(r"^\d+[A-Z]\d+$")


def copy_cell_style(style_data, target_cell):
    """Copy format từ dict dữ liệu tạm sang cell đích."""
    if style_data.get("font") is not None:
        target_cell.font = copy(style_data["font"])
    if style_data.get("border") is not None:
        target_cell.border = copy(style_data["border"])
    if style_data.get("fill") is not None:
        target_cell.fill = copy(style_data["fill"])
    if style_data.get("alignment") is not None:
        target_cell.alignment = copy(style_data["alignment"])
    if style_data.get("number_format") is not None:
        target_cell.number_format = style_data["number_format"]


def extract_cell_payload(cell):
    """Lưu value + style của cell dưới dạng dict để tránh lỗi sai kiểu dữ liệu."""
    return {
        "value": cell.value,
        "font": copy(cell.font) if cell.font is not None else None,
        "border": copy(cell.border) if cell.border is not None else None,
        "fill": copy(cell.fill) if cell.fill is not None else None,
        "number_format": cell.number_format,
        "alignment": copy(cell.alignment) if cell.alignment is not None else None,
    }


def to_range_string(min_col, min_row, max_col, max_row):
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"


def get_merged_boundaries(ws):
    return [range_boundaries(str(rng)) for rng in ws.merged_cells.ranges]


def remap_merged_ranges_for_insert(merged_ranges, start_row, amount):
    remapped = []
    for min_col, min_row, max_col, max_row in merged_ranges:
        if max_row < start_row:
            remapped.append((min_col, min_row, max_col, max_row))
        elif min_row >= start_row:
            remapped.append((min_col, min_row + amount, max_col, max_row + amount))
        else:
            # Range cắt ngang vị trí chèn cần được mở rộng xuống.
            remapped.append((min_col, min_row, max_col, max_row + amount))
    return remapped


def remap_merged_ranges_for_delete(merged_ranges, start_row, amount):
    end_row = start_row + amount - 1
    remapped = []

    for min_col, min_row, max_col, max_row in merged_ranges:
        if max_row < start_row:
            remapped.append((min_col, min_row, max_col, max_row))
            continue

        if min_row > end_row:
            remapped.append((min_col, min_row - amount, max_col, max_row - amount))
            continue

        # Merge nằm hoàn toàn trong vùng bị xóa thì bỏ luôn.
        if min_row >= start_row and max_row <= end_row:
            continue

        # Merge phủ qua toàn vùng xóa, cần co ngắn lại.
        if min_row < start_row and max_row > end_row:
            remapped.append((min_col, min_row, max_col, max_row - amount))
            continue

        # Merge nằm phía trên và đè xuống vùng xóa.
        if min_row < start_row <= max_row <= end_row:
            new_max_row = start_row - 1
            if new_max_row >= min_row:
                remapped.append((min_col, min_row, max_col, new_max_row))
            continue

        # Merge bắt đầu trong vùng xóa và kéo xuống dưới vùng xóa.
        if start_row <= min_row <= end_row < max_row:
            new_min_row = start_row
            new_max_row = max_row - amount
            if new_max_row >= new_min_row:
                remapped.append((min_col, new_min_row, max_col, new_max_row))

    return remapped


def apply_merged_ranges(ws, merged_ranges):
    ws._cells = {
        coord: cell
        for coord, cell in ws._cells.items()
        if not isinstance(cell, MergedCell)
    }
    ws.merged_cells = MultiCellRange()

    seen = set()
    for min_col, min_row, max_col, max_row in merged_ranges:
        if min_row > max_row or min_col > max_col:
            continue
        range_str = to_range_string(min_col, min_row, max_col, max_row)
        if range_str in seen:
            continue
        ws.merge_cells(range_str)
        seen.add(range_str)


def snapshot_row_heights(ws, start_row):
    return {
        row_idx: dim.height
        for row_idx, dim in ws.row_dimensions.items()
        if row_idx >= start_row and dim.height is not None
    }


def shift_row_heights_after_insert(ws, row_heights, start_row, amount):
    for row_idx in sorted(row_heights.keys(), reverse=True):
        ws.row_dimensions[row_idx + amount].height = row_heights[row_idx]

    for row_idx in range(start_row, start_row + amount):
        ws.row_dimensions[row_idx].height = None


def process_excel(uploaded_file, class_name, filename):
    wb = None
    try:
        source_start_row = 20
        cut_amount = 8
        source_end_row = source_start_row + cut_amount - 1
        target_insert_row = 2

        # Load workbook hoàn toàn trong RAM
        file_buffer = io.BytesIO(uploaded_file.getvalue())
        wb = openpyxl.load_workbook(file_buffer)
        
        # Khớp động tên sheet theo tên lớp user nhập
        src_sheet_name = f"MH {class_name} (2025-2026)"
        tgt_sheet_name = f"NL, PC {class_name} (2025-2026)"
        
        if src_sheet_name not in wb.sheetnames or tgt_sheet_name not in wb.sheetnames:
            missing_sheets = [
                sheet_name
                for sheet_name in (src_sheet_name, tgt_sheet_name)
                if sheet_name not in wb.sheetnames
            ]
            return None, f"❌ File '{filename}' thiếu sheet: {', '.join(missing_sheets)}."
            
        ws_src = wb[src_sheet_name]
        ws_tgt = wb[tgt_sheet_name]

        src_merged_ranges = get_merged_boundaries(ws_src)
        cut_block_merged_ranges = []
        for min_col, min_row, max_col, max_row in src_merged_ranges:
            overlaps = not (max_row < source_start_row or min_row > source_end_row)
            fully_inside = min_row >= source_start_row and max_row <= source_end_row

            if overlaps and not fully_inside:
                range_str = to_range_string(min_col, min_row, max_col, max_row)
                return (
                    None,
                    f"❌ File '{filename}' có merge cell cắt ngang vùng dòng 20-27 ({range_str}), không thể cắt an toàn.",
                )

            if fully_inside:
                new_min_row = target_insert_row + (min_row - source_start_row)
                new_max_row = target_insert_row + (max_row - source_start_row)
                cut_block_merged_ranges.append((min_col, new_min_row, max_col, new_max_row))
        
        # 1. Lưu lại data và format của dòng 20 đến 27
        rows_data = []
        source_row_heights = []
        for row_idx, row in enumerate(
            ws_src.iter_rows(
                min_row=source_start_row,
                max_row=source_end_row,
                max_col=ws_src.max_column,
            ),
            start=source_start_row,
        ):
            row_data = [extract_cell_payload(cell) for cell in row]
            rows_data.append(row_data)
            source_row_heights.append(ws_src.row_dimensions[row_idx].height)

        # 2. Insert 8 dòng vào sheet đích và remap merge/row-height để giữ layout.
        tgt_merged_ranges_before_insert = get_merged_boundaries(ws_tgt)
        tgt_row_heights_before_insert = snapshot_row_heights(ws_tgt, start_row=target_insert_row)

        ws_tgt.insert_rows(target_insert_row, amount=cut_amount)

        remapped_tgt_merged_ranges = remap_merged_ranges_for_insert(
            tgt_merged_ranges_before_insert,
            start_row=target_insert_row,
            amount=cut_amount,
        )
        apply_merged_ranges(ws_tgt, remapped_tgt_merged_ranges)
        shift_row_heights_after_insert(
            ws_tgt,
            tgt_row_heights_before_insert,
            start_row=target_insert_row,
            amount=cut_amount,
        )
        
        # 3. Paste dữ liệu và format vào sheet đích (từ dòng 2 đến 9)
        for r_offset, row_data in enumerate(rows_data):
            target_row_idx = target_insert_row + r_offset
            row_height = source_row_heights[r_offset]
            if row_height is not None:
                ws_tgt.row_dimensions[target_row_idx].height = row_height

            for c_idx, cell_data in enumerate(row_data, start=1):
                tgt_cell = ws_tgt.cell(row=target_row_idx, column=c_idx)
                tgt_cell.value = cell_data['value']
                copy_cell_style(cell_data, tgt_cell)

        # Áp merge từ block cắt sang vùng mới dán ở sheet đích.
        for min_col, min_row, max_col, max_row in cut_block_merged_ranges:
            ws_tgt.merge_cells(to_range_string(min_col, min_row, max_col, max_row))
                
        # 4. Cắt (Xóa) dòng 20-27 ở sheet nguồn và remap merge còn lại.
        ws_src.delete_rows(source_start_row, amount=cut_amount)
        remapped_src_merged_ranges = remap_merged_ranges_for_delete(
            src_merged_ranges,
            start_row=source_start_row,
            amount=cut_amount,
        )
        apply_merged_ranges(ws_src, remapped_src_merged_ranges)
        
        # Xuất file đã xử lý ra memory (BytesIO)
        out_stream = io.BytesIO()
        wb.save(out_stream)
        out_stream.seek(0)
        
        return out_stream, None
    except Exception as e:
        return None, f"❌ Lỗi khi xử lý file '{filename}': {str(e)}"
    finally:
        if wb is not None:
            wb.close()

# Giao diện chính
st.title("Chỉnh sửa học bạ Excel")

# 1. Nhập tên lớp với Validation (bắt buộc dạng xAy)
class_input = st.text_input("Nhập tên lớp (Định dạng xAy, VD: 4A1, 5A2):").strip().upper()

# 2. Upload hàng loạt file
uploaded_files = st.file_uploader("Kéo thả nhiều file Excel (.xlsx) vào đây", type=["xlsx"], accept_multiple_files=True)

if st.button("🚀 Bắt đầu xử lý", type="primary"):
    # Ràng buộc điều kiện
    if not class_input:
        st.warning("⚠️ Vui lòng nhập tên lớp.")
    elif not CLASS_NAME_PATTERN.fullmatch(class_input):
        st.error("⚠️ Tên lớp sai định dạng! Vui lòng nhập đúng chuẩn xAy (VD: 4A1).")
    elif not uploaded_files:
        st.warning("⚠️ Vui lòng tải lên ít nhất 1 file Excel.")
    else:
        with st.spinner("Đang xử lý..."):
            processed_files = []
            errors = []
            
            for file in uploaded_files:
                out_stream, err_msg = process_excel(file, class_input, file.name)
                if err_msg:
                    errors.append(err_msg)
                else:
                    processed_files.append((f"{file.name}", out_stream))
            
            # Hiển thị lỗi nếu có
            for error in errors:
                st.error(error)
                
            # Trả file cho người dùng tải về
            if processed_files:
                st.success(f"✅ Đã xử lý thành công {len(processed_files)} file!")
                
                # Nếu có nhiều file, nén thành file ZIP cho tiện
                if len(processed_files) > 1:
                    zip_buffer = io.BytesIO()
                    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                        for file_name, file_stream in processed_files:
                            zip_file.writestr(file_name, file_stream.getvalue())
                    
                    st.download_button(
                        label="📦 Tải tất cả file đã xử lý (.zip)",
                        data=zip_buffer.getvalue(),
                        file_name=f"Processed_Excel_Files_{class_input}.zip",
                        mime="application/zip"
                    )
                else:
                    # Nếu chỉ có 1 file, tải trực tiếp file .xlsx
                    file_name, file_stream = processed_files[0]
                    st.download_button(
                        label="💾 Tải file đã xử lý (.xlsx)",
                        data=file_stream.getvalue(),
                        file_name=file_name,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )